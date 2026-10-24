"""
Excel Exporter for MRP Procurement Plans.
Generates production-grade spreadsheets (.xlsx) for Natura Siberica Supply Chain & Finance.
"""

from pathlib import Path
from typing import Union
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import ProcurementPlan


class MRPExcelExporter:
    """
    Renders structured ProcurementPlan models into styled Excel workbooks.
    """

    @classmethod
    def export_to_file(cls, plan: ProcurementPlan, output_path: Union[str, Path]) -> str:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "План закупки сырья"
        ws.views.sheetView[0].showGridLines = True

        # Styles
        title_font = Font(name="Calibri", size=15, bold=True, color="1F4E79")
        subtitle_font = Font(name="Calibri", size=10, italic=True, color="595959")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        total_font = Font(name="Calibri", size=11, bold=True, color="000000")
        total_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9")
        )

        # 1. Header Block
        ws.merge_cells("A1:K1")
        ws["A1"] = "NATURA SIBERICA | ПЛАН МАТЕРИАЛЬНО-ТЕХНИЧЕСКОГО СНАБЖЕНИЯ (MRP RUN)"
        ws["A1"].font = title_font
        ws["A1"].alignment = Alignment(vertical="center")

        ws.merge_cells("A2:K2")
        ws["A2"] = f"Продукция: {plan.product_name} (SKU: {plan.sku}) | Сформировано: {plan.generated_at}"
        ws["A2"].font = subtitle_font

        # Summary KPIs
        ws["A4"] = "Целевой выпуск партии:"
        ws["B4"] = f"{plan.target_production_volume:,} шт."
        ws["B4"].font = Font(bold=True)

        ws["D4"] = "Критический срок поставки (Lead Time):"
        ws["E4"] = f"{plan.critical_lead_time_days} дн."
        ws["E4"].font = Font(bold=True)

        ws["G4"] = "ИТОГО БЮДЖЕТ ЗАКУПКИ:"
        ws["H4"] = f"{plan.total_procurement_budget_rub:,.2f} руб."
        ws["H4"].font = Font(bold=True, color="C00000")

        # 2. Table Headers
        headers = [
            ("A6", "Код InvenTree"),
            ("B6", "Наименование сырья / упаковки"),
            ("C6", "Ед. изм."),
            ("D6", "Валовая потр."),
            ("E6", "Страх. запас"),
            ("F6", "Свободный остаток"),
            ("G6", "В пути"),
            ("H6", "Чистая потр."),
            ("I6", "К заказу (MOQ)"),
            ("J6", "Сумма (руб.)"),
            ("K6", "Статус снабжения")
        ]

        for cell_ref, text in headers:
            cell = ws[cell_ref]
            cell.value = text
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # 3. Table Rows
        start_row = 7
        for idx, item in enumerate(plan.requirements):
            row = start_row + idx
            ws[f"A{row}"] = item.item_code
            ws[f"B{row}"] = item.name
            ws[f"C{row}"] = item.unit
            ws[f"D{row}"] = item.gross_requirement
            ws[f"E{row}"] = item.safety_stock
            ws[f"F{row}"] = item.current_free_stock
            ws[f"G{row}"] = item.in_transit_stock
            ws[f"H{row}"] = item.net_requirement
            ws[f"I{row}"] = item.recommended_order_qty
            ws[f"J{row}"] = item.total_order_cost_rub
            ws[f"K{row}"] = item.urgency_status

            # Alignment and number formats
            ws[f"A{row}"].alignment = Alignment(horizontal="center")
            ws[f"C{row}"].alignment = Alignment(horizontal="center")
            ws[f"D{row}"].number_format = "#,##0.0"
            ws[f"E{row}"].number_format = "#,##0.0"
            ws[f"F{row}"].number_format = "#,##0.0"
            ws[f"G{row}"].number_format = "#,##0.0"
            ws[f"H{row}"].number_format = "#,##0.0"
            ws[f"I{row}"].number_format = "#,##0.0"
            ws[f"J{row}"].number_format = "#,##0.00"

            for col_letter in ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K"]:
                ws[f"{col_letter}{row}"].border = thin_border

        # 4. Total Row
        total_row = start_row + len(plan.requirements)
        ws.merge_cells(f"A{total_row}:I{total_row}")
        ws[f"A{total_row}"] = "ИТОГО ПО СПЕЦИФИКАЦИИ:"
        ws[f"A{total_row}"].font = total_font
        ws[f"A{total_row}"].alignment = Alignment(horizontal="right")
        ws[f"A{total_row}"].fill = total_fill

        ws[f"J{total_row}"] = f"=SUM(J{start_row}:J{total_row - 1})"
        ws[f"J{total_row}"].font = total_font
        ws[f"J{total_row}"].number_format = "#,##0.00"
        ws[f"J{total_row}"].fill = total_fill

        ws[f"K{total_row}"].fill = total_fill

        for col_letter in ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K"]:
            ws[f"{col_letter}{total_row}"].border = thin_border

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        ws.column_dimensions["B"].width = 38
        ws.column_dimensions["K"].width = 36

        # Save
        out_p = Path(output_path)
        out_p.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(out_p))
        return str(out_p)
