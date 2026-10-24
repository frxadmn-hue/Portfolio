"""
Word (.docx) Exporter for Legal Protocol of Disagreements.
Generates ready-to-sign Russian commercial legal documents from AuditReport.
"""

from pathlib import Path
from typing import Union
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from .models import AuditReport


class LegalDocxExporter:
    """
    Renders AuditReport protocol items into an official corporate Microsoft Word (.docx) document.
    """

    @classmethod
    def export_protocol_to_docx(cls, report: AuditReport, output_path: Union[str, Path]) -> str:
        doc = docx.Document()

        # Set 2 cm margins
        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        # 1. Document Title
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_title = p_title.add_run("ПРОТОКОЛ РАЗНОГЛАСИЙ\n")
        run_title.font.name = "Times New Roman"
        run_title.font.size = Pt(14)
        run_title.font.bold = True

        run_sub = p_title.add_run(f"к проекту: {report.contract_title}")
        run_sub.font.name = "Times New Roman"
        run_sub.font.size = Pt(12)
        run_sub.font.bold = True

        # Date & City line
        p_city = doc.add_paragraph()
        p_city.paragraph_format.space_after = Pt(14)
        r_city = p_city.add_run("г. Москва\t\t\t\t\t\t\t«_____» ____________ 2026 г.")
        r_city.font.name = "Times New Roman"
        r_city.font.size = Pt(11)

        # Preamble
        p_pre = doc.add_paragraph()
        p_pre.paragraph_format.space_after = Pt(12)
        p_pre.paragraph_format.line_spacing = 1.15
        r_pre = p_pre.add_run(
            "ООО «Натура Сиберика», именуемое в дальнейшем «Поставщик», с одной стороны, и "
            "Торговая сеть, именуемая в дальнейшем «Покупатель», с другой стороны, "
            "составили настоящий Протокол разногласий о нижеследующем:"
        )
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11)

        # 2. Main 3-Column Table
        table = doc.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # Set column widths
        widths = [Inches(1.0), Inches(2.8), Inches(3.2)]
        for i, w in enumerate(widths):
            table.columns[i].width = w

        # Header Row
        hdr_cells = table.rows[0].cells
        hdr_titles = ["Пункт", "Редакция Покупателя (Сети)", "Редакция Поставщика (Natura Siberica) и правовое обоснование"]
        
        for idx, title in enumerate(hdr_titles):
            cell = hdr_cells[idx]
            cell.width = widths[idx]
            cell.text = title
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in cell.paragraphs[0].runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(10)
                r.font.bold = True
            # Background fill
            shading = parse_xml(r'<w:shd {} w:fill="F2F2F2"/>'.format(nsdecls('w')))
            cell._tc.get_or_add_tcPr().append(shading)

        # Data Rows
        for item in report.protocol_items:
            row_cells = table.add_row().cells
            for idx, w in enumerate(widths):
                row_cells[idx].width = w

            # Column 1: Clause
            p0 = row_cells[0].paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r0 = p0.add_run(f"п. {item.contract_clause_ref}")
            r0.font.name = "Times New Roman"
            r0.font.size = Pt(10)
            r0.font.bold = True

            # Column 2: Buyer wording
            p1 = row_cells[1].paragraphs[0]
            r1 = p1.add_run(item.retailer_wording)
            r1.font.name = "Times New Roman"
            r1.font.size = Pt(9.5)

            # Column 3: Supplier wording + Legal rationale
            p2 = row_cells[2].paragraphs[0]
            r2 = p2.add_run(item.supplier_counter_wording + "\n\n")
            r2.font.name = "Times New Roman"
            r2.font.size = Pt(9.5)
            r2.font.bold = True

            r2_rat = p2.add_run(f"Обоснование: {item.legal_rationale}")
            r2_rat.font.name = "Times New Roman"
            r2_rat.font.size = Pt(8.5)
            r2_rat.font.italic = True
            r2_rat.font.color.rgb = RGBColor(89, 89, 89)

        # Set borders for table
        tblPr = table._tbl.tblPr
        tblBorders = parse_xml(
            r'<w:tblBorders {} >'
            r'<w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            r'<w:left w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            r'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            r'<w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
            r'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
            r'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
            r'</w:tblBorders>'.format(nsdecls('w'))
        )
        tblPr.append(tblBorders)

        # 3. Signatures
        doc.add_paragraph().paragraph_format.space_before = Pt(20)
        p_sig = doc.add_paragraph()
        p_sig.paragraph_format.line_spacing = 1.3
        
        sig_text = (
            "ПОДПИСИ И РЕКВИЗИТЫ СТОРОН:\n\n"
            "От Поставщика:                                              От Покупателя:\n"
            "ООО «Натура Сиберика»                                       Торговая сеть / Ритейлер\n\n"
            "_________________ / _________________ /                      _________________ / _________________ /\n"
            "М.П.                                                        М.П.\n"
        )
        r_sig = p_sig.add_run(sig_text)
        r_sig.font.name = "Times New Roman"
        r_sig.font.size = Pt(10)

        # Save
        out_p = Path(output_path)
        out_p.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out_p))
        return str(out_p)
